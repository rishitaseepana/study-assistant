from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)

class Exporter:

    def __init__(self):
        self.generated_time = datetime.now().strftime(
            "%d %B %Y %I:%M %p"
        )

    def export(
        self,
        text,
        filename,
        file_type,
        title="AI Study Assistant",
        topic=None
    ):

        file_type = file_type.lower()
        if file_type == "txt":
            return self.export_txt(
                text,
                filename,
                title,
                topic
            )
        if file_type == "docx":
            return self.export_docx(
                text,
                filename,
                title,
                topic
            )
        if file_type == "pdf":
            return self.export_pdf(
                text,
                filename,
                title,
                topic
            )

        raise ValueError(
            f"Unsupported format: {file_type}"
        )

    def export_temp(
        self,
        text,
        file_type,
        title="AI Study Assistant",
        topic=None
    ):

        with NamedTemporaryFile(
            delete=False,
            suffix=f".{file_type}"
        ) as tmp:

            filename = Path(tmp.name).with_suffix("")

        return self.export(
            text=text,
            filename=filename,
            file_type=file_type,
            title=title,
            topic=topic
        )

    def export_txt(
        self,
        text,
        filename,
        title,
        topic
    ):

        path = Path(filename).with_suffix(".txt")

        with open(path, "w", encoding="utf-8") as f:
            f.write("=" * 60 + "\n")
            f.write(title.upper() + "\n")
            f.write("=" * 60 + "\n\n")
            if topic:
                f.write(f"Topic : {topic}\n")
            f.write(f"Generated : {self.generated_time}\n")
            f.write("\n" + "=" * 60 + "\n\n")
            f.write(text)

        return str(path)

    def export_docx(
        self,
        text,
        filename,
        title,
        topic
    ):

        path = Path(filename).with_suffix(".docx")
        document = Document()
        section = document.sections[0]
        section.top_margin = Pt(50)
        section.bottom_margin = Pt(50)
        section.left_margin = Pt(50)
        section.right_margin = Pt(50)
        heading = document.add_heading(title, level=0)
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        info = document.add_paragraph()

        if topic:
            info.add_run("Topic: ").bold = True
            info.add_run(topic + "\n")

        info.add_run("Generated: ").bold = True
        info.add_run(self.generated_time)
        document.add_paragraph()
        self._write_docx_content(
            document,
            text
        )
        document.save(path)

        return str(path)

    def export_pdf(
        self,
        text,
        filename,
        title,
        topic
    ):

        path = Path(filename).with_suffix(".pdf")
        styles = getSampleStyleSheet()
        title_style = styles["Title"]
        title_style.alignment = TA_CENTER
        heading_style = styles["Heading2"]
        body_style = styles["BodyText"]
        story = []
        story.append(
            Paragraph(title, title_style)
        )
        story.append(Spacer(1, 20))

        if topic:
            story.append(
                Paragraph(
                    f"<b>Topic:</b> {topic}",
                    body_style
                )
            )

        story.append(
            Paragraph(
                f"<b>Generated:</b> {self.generated_time}",
                body_style
            )
        )

        story.append(
            Spacer(1, 20)
        )

        for line in text.splitlines():
            line = line.strip()

            if not line:
                story.append(Spacer(1, 10))
                continue

            if line.startswith("# "):
                story.append(
                    Paragraph(
                        line[2:],
                        heading_style
                    )
                )

            elif line.startswith("## "):
                story.append(
                    Paragraph(
                        f"<b>{line[3:]}</b>",
                        body_style
                    )
                )

            elif line.startswith("- "):
                story.append(
                    Paragraph(
                        f"• {line[2:]}",
                        body_style
                    )
                )

            elif (
                len(line) > 2 and
                line[0].isdigit() and
                line[1] == "."
            ):
                story.append(
                    Paragraph(
                        line,
                        body_style
                    )
                )

            else:
                story.append(
                    Paragraph(
                        line,
                        body_style
                    )
                )

        pdf = SimpleDocTemplate(
            str(path)
        )

        pdf.build(
            story,
            onFirstPage=self._add_page_number,
            onLaterPages=self._add_page_number
        )

        return str(path)

    def _write_docx_content(
        self,
        document,
        text
    ):

        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            if line.startswith("# "):
                document.add_heading(
                    line[2:],
                    level=1
                )

            elif line.startswith("## "):
                document.add_heading(
                    line[3:],
                    level=2
                )

            elif line.startswith("- "):
                document.add_paragraph(
                    line[2:],
                    style="List Bullet"
                )

            elif (
                len(line) > 2 and
                line[0].isdigit() and
                line[1] == "."
            ):
                document.add_paragraph(
                    line,
                    style="List Number"
                )

            else:
                p = document.add_paragraph()
                p.add_run(line)

    def _add_page_number(
        self,
        canvas,
        doc
    ):

        page = canvas.getPageNumber()
        canvas.setFont(
            "Helvetica",
            9
        )
        canvas.drawRightString(
            550,
            20,
            f"Page {page}"
        )